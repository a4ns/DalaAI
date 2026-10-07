"""Actual report PostgreSQL gates: frozen source, restricted LOGIN and app.main."""
import os
from pathlib import Path
import sys
import unittest
from unittest.mock import patch

ROOT=Path(__file__).resolve().parents[1]
if not os.environ.get('DALA_TEST_DATABASE_URL'):
    raise SystemExit('NOT_RUN: explicit disposable PostgreSQL database required')
os.environ['DALA_C_RUNTIME_POSTGRES_ACCEPTANCE']='1'
sys.path[:0]=[str(ROOT/'backend'),str(ROOT/'backend/tests'),
             str(ROOT/'backend/tests/integration'),str(ROOT/'ops/provision')]
from test_persistence_role import ApplicationRoleTests
from test_c_runtime_reports_postgres import RuntimeReportsPostgresTests, SESSION_COOKIE_NAME
from database_profile import grant_existing_runtime_role

class RestrictedReports(RuntimeReportsPostgresTests):
    def runtime_connect(self):
        return self.pg.connect(ApplicationRoleTests.runtime_dsn,autocommit=True,connect_timeout=3,
            options=f'-c search_path={self.schema} -c statement_timeout=10000 -c lock_timeout=3000',
            row_factory=self.dict_row)

    def setUp(self):
        super().setUp()
        with self.connect() as owner:
            owner.execute((ROOT/'backend/db/proposals/003_auth_rate_limits.sql').read_text())
        grant_existing_runtime_role(self.connect,self.runtime_connect,self.schema)
        self.sessions.connect=self.runtime_connect
        with self.runtime_connect() as db:
            row=db.execute('SELECT current_user AS role,session_user AS login').fetchone()
        self.assertEqual(row['role'],ApplicationRoleTests.runtime_role)
        self.assertEqual(row['login'],row['role'])

class MountedReports(RestrictedReports):
    def test_actual_main_mounts_all_three_protected_routes(self):
        from app.main import create_app
        from app.runtime import RuntimeSettings
        from fastapi.testclient import TestClient
        origin='https://reports.test'
        settings=RuntimeSettings(mode='demo',database_url=ApplicationRoleTests.runtime_dsn,
            allowed_origin=origin,database_schema=self.schema)
        # Synthetic clock fixes fixture instants only; real PostgreSQL/auth and
        # actual app.main lifespan/role checks remain in use.
        with patch('app.core.auth_boundary.SystemRealClock',return_value=self.real):
            app=create_app(settings=settings,connect=self.runtime_connect)
        with TestClient(app,base_url=origin,client=('127.0.0.1',45000)) as client:
            self.assertEqual(client.get('/readyz').status_code,200)
            for route in ('/analytics/shift','/reports/shift','/reports/orders/'+self.order):
                self.assertEqual(client.get('/api/v1'+route,params=self.query).status_code,401)
                response=client.get('/api/v1'+route,params=self.query,
                    headers={'cookie':SESSION_COOKIE_NAME+'='+self.handle})
                self.assertEqual(response.status_code,200,'Mounted report route failed')
                self.assertIn('no-store',response.headers['cache-control'])


    def test_actual_main_binary_exports_and_access_errors(self):
        from app.main import create_app
        from app.runtime import RuntimeSettings
        from fastapi.testclient import TestClient
        from io import BytesIO
        from openpyxl import load_workbook
        from hashlib import sha256
        from uuid import uuid4
        from datetime import timedelta
        origin='https://reports.test'
        settings=RuntimeSettings(mode='demo',database_url=ApplicationRoleTests.runtime_dsn,
            allowed_origin=origin,database_schema=self.schema)
        executor_handle=uuid4().hex
        with self.connect() as owner:
            owner.execute("""INSERT INTO auth_sessions(id,employee_id,token_hash,csrf_token,created_at,expires_at)
                VALUES (%s,%s,%s,%s,%s,%s)""", (str(uuid4()),self.executor,
                sha256(executor_handle.encode()).hexdigest(),uuid4().hex,
                self.real.now()-timedelta(minutes=1),self.real.now()+timedelta(hours=1)))
            before={table:owner.execute('SELECT count(*) AS n FROM '+table).fetchone()['n']
                    for table in ('orders','submissions','reviews','order_events','auth_sessions','delivery_jobs')}
        with patch('app.core.auth_boundary.SystemRealClock',return_value=self.real):
            app=create_app(settings=settings,connect=self.runtime_connect)
        with TestClient(app,base_url=origin,client=('127.0.0.1',45000)) as client:
            self.assertEqual(client.get('/readyz').status_code,200)
            for suffix in ('pdf','xlsx'):
                for stem in ('shift','orders/'+self.order):
                    path='/api/v1/reports/'+stem+'.'+suffix
                    self.assertEqual(client.get(path,params=self.query).status_code,401)
                    self.assertEqual(client.get(path,params=self.query,
                        headers={'cookie':SESSION_COOKIE_NAME+'='+executor_handle}).status_code,403)
                    response=client.get(path,params=self.query,
                        headers={'cookie':SESSION_COOKIE_NAME+'='+self.handle})
                    self.assertEqual(response.status_code,200,'Mounted binary report failed')
                    self.assertIn('no-store',response.headers['cache-control'])
                    self.assertIn('attachment;',response.headers['content-disposition'])
                    if suffix=='pdf':
                        self.assertTrue(response.content.startswith(b'%PDF-'))
                        self.assertIn(b'DejaVuSans',response.content)
                    else:
                        book=load_workbook(BytesIO(response.content),read_only=True,data_only=False)
                        self.assertTrue(book.sheetnames)
                        self.assertTrue(any(cell.value=='Synthetic' for sheet in book for row in sheet for cell in row)
                                        if stem.startswith('orders/') else True)
                        self.assertFalse(any(cell.data_type=='f' for sheet in book for row in sheet for cell in row))
                        book.close()
                    invalid=client.get(path,params={**self.query,'format':'html'},
                        headers={'cookie':SESSION_COOKIE_NAME+'='+self.handle})
                    self.assertEqual(invalid.status_code,422)
                    self.assertNotIn('content-disposition',invalid.headers)
        with self.connect() as owner:
            after={table:owner.execute('SELECT count(*) AS n FROM '+table).fetchone()['n'] for table in before}
        self.assertEqual(before,after)

def run(suite,expected,label):
    assert suite.countTestCases()==expected,label+' count changed'
    result=unittest.TextTestRunner(verbosity=2).run(suite)
    if result.skipped or result.testsRun!=expected or not result.wasSuccessful():
        raise SystemExit('FAIL: '+label)
    print(f'PASS: {label}, {expected} actual PostgreSQL tests, zero skips')

run(unittest.defaultTestLoader.loadTestsFromTestCase(RuntimeReportsPostgresTests),8,'C111 author reports')
ApplicationRoleTests.setUpClass()
try:
    run(unittest.defaultTestLoader.loadTestsFromTestCase(RestrictedReports),8,'C111 restricted API LOGIN')
    run(unittest.TestSuite([MountedReports('test_actual_main_mounts_all_three_protected_routes')]),1,'Actual app.main reports')
    run(unittest.TestSuite([MountedReports('test_actual_main_binary_exports_and_access_errors')]),1,'Actual app.main PDF and XLSX exports')
finally:
    ApplicationRoleTests.doClassCleanups()
    if ApplicationRoleTests.tearDown_exceptions:
        raise SystemExit('FAIL: report-role cleanup not confirmed')
