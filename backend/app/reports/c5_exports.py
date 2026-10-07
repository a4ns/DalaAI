"""Bounded PDF/XLSX rendering from authorized C3 facts, with no source I/O.

The HTTP adapter invokes this inside RuntimeReportService.capture(project=...),
before its final real-clock/session recheck. Never deserialize facts from HTTP.
Fonts are fixed local assets. No URL fetches, images, formulas, or hyperlinks.
"""
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from io import BytesIO
import math
import os
import pickle
import subprocess
import sys
from pathlib import Path
import time
from zoneinfo import ZoneInfo

from app.analytics.c3_types import AnalyticsFacts
from app.orders.models import DomainError

MAX_BYTES = 8 * 1024 * 1024
MAX_CELLS = 120_000
MAX_TEXT_CHARS = 1_000_000
MAX_CELL_CHARS = 30_000  # Below Excel's 32,767-character silent-truncation boundary.
MAX_PDF_PAGES = 200
RENDER_SECONDS = 6.0
FONT_REGULAR = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
FONT_BOLD = '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'
ZONE = ZoneInfo('Asia/Almaty')

STATUS = {'issued': 'Выдан', 'queued': 'В очереди', 'accepted': 'Принят',
          'rejected': 'Отклонён', 'in_progress': 'В работе', 'paused': 'Приостановлен',
          'done': 'Результат отправлен', 'ai_review': 'На проверке',
          'rework': 'На доработке', 'closed': 'Закрыт', 'cancelled': 'Отменён'}
METRICS = {'issued_orders': 'Выдано нарядов', 'submitted_orders': 'Отправлено нарядов',
           'submission_attempts': 'Отправлено попыток', 'closed_orders': 'Закрыто нарядов',
           'rework_decisions': 'Решений о доработке', 'awaiting_review': 'Ожидают проверки',
           'overdue_active': 'Просрочено на время снимка', 'attempt_on_time': 'Попытки в срок',
           'human_score': 'Оценка мастера среди закрытых',
           'human_score_mean': 'Средняя оценка мастера среди закрытых',
           'closed_on_time': 'В срок среди закрытых', 'closed_with_rework': 'Закрытые с доработкой'}
METRIC_STATUS = {'ok': 'Доступно', 'missing': 'Нет данных', 'partial': 'Частичные данные',
                 'no_cohort': 'Нет подходящих записей'}
MODE = {'model': 'Модель', 'rules_fallback': 'Правила (резервный режим)', 'manual': 'Ручной режим'}
RECOMMENDATION = {'satisfactory': 'Удовлетворительно', 'rework_recommended': 'Рекомендуется доработка',
                  'needs_master_review': 'Требуется проверка мастера'}
LIMITATIONS = (
    'Показана только разрешённая область источника, не итог предприятия.',
    'Оценка ИИ — рекомендация. Производственное решение принимает мастер.',
    'Фото: только идентификаторы. Изображения не встроены и не проверялись этим отчётом.',
    'История попыток и решений не является полным журналом событий наряда.',
    'Материалы заявлены исполнителями; это не подтверждённые складские списания.',
    'Пустая ячейка означает отсутствие значения, не ноль. Ноль сохранён как число.',
    'Время отображено в Asia/Almaty (UTC+05:00). Период [начало, конец).',
    'Числа Excel ограничены его точностью; точные десятичные значения даны в соседних текстовых полях.',
    'Длинный текст Excel сохранён целиком; для просмотра откройте ячейку или строку формул.',
)


def _limit():
    return DomainError('REPORT_LIMIT_EXCEEDED', 'Report export exceeds rendering limits')


def _unavailable():
    return DomainError('TEMPORARILY_UNAVAILABLE', 'Report export is unavailable')


class _Budget:
    def __init__(self):
        self.cells = self.chars = 0
        self.deadline = time.monotonic() + RENDER_SECONDS

    def check(self):
        if time.monotonic() > self.deadline:
            raise _unavailable()

    def row(self, values):
        self.check()
        self.cells += len(values)
        if self.cells > MAX_CELLS:
            raise _limit()
        for value in values:
            if isinstance(value, str):
                # Fail instead of silently dropping XML-invalid control characters.
                if (len(value) > MAX_CELL_CHARS or any(
                        not (c in '\t\n\r' or 0x20 <= ord(c) <= 0xD7FF
                             or 0xE000 <= ord(c) <= 0xFFFD or 0x10000 <= ord(c) <= 0x10FFFF)
                        for c in value)):
                    raise _limit()
                self.chars += len(value)
            elif isinstance(value, Decimal):
                if (not value.is_finite() or not math.isfinite(float(value))
                        or value != 0 and float(value) == 0):
                    raise _unavailable()
            elif isinstance(value, datetime) and (value.tzinfo is None or value.utcoffset() is None):
                raise _unavailable()
        if self.chars > MAX_TEXT_CHARS:
            raise _limit()


class _BoundedBytes(BytesIO):
    _failed = False

    def write(self, data):
        if self._failed:
            # ZipFile finalization after the first failure must not produce an
            # unraisable destructor error. These bytes can never be returned.
            return len(data)
        if self.tell() + len(data) > MAX_BYTES:
            self._failed = True
            raise _limit()
        return super().write(data)


@dataclass
class _Table:
    title: str
    headers: tuple
    rows: list
    budget: _Budget

    def add(self, *values):
        if len(values) != len(self.headers):
            raise ValueError('Export row shape mismatch')
        self.budget.row(values)
        self.rows.append(values)


def _tables(facts, order_id, historical_evidence, budget):
    if not isinstance(facts, AnalyticsFacts):
        raise TypeError('C5 requires AnalyticsFacts')
    tables = []

    def table(title, *headers):
        budget.row(headers)
        result = _Table(title, headers, [], budget)
        tables.append(result)
        return result

    selected = None
    if order_id is not None:
        choices = [row for row in facts.orders if row.order.id == order_id]
        if len(choices) != 1:
            raise ValueError('Order is not uniquely present in authorized facts')
        selected = choices[0]
    title = 'Отчёт смены' if selected is None else 'Наряд ' + str(selected.order.number)
    p = facts.provenance
    meta = table('Об отчёте', 'Поле', 'Значение')
    for label, value in (
        ('Отчёт', title), ('Режим', 'СИНТЕТИЧЕСКИЕ ДАННЫЕ' if p.synthetic else 'ДАННЫЕ ИСТОЧНИКА: НЕ СИНТЕТИЧЕСКИЕ'),
        ('Источник', p.source_ref), ('Разрешённая область', p.scope_description),
        ('Версия фактов', facts.schema_version), ('Полнота выборки', p.coverage),
        ('История полная по заявлению источника', p.history_complete),
        ('Доменное время снимка', p.domain_as_of), ('Реальное время получения', p.captured_at_real),
        ('Начало периода, включено', facts.period.start), ('Конец периода, исключён', facts.period.end),
        ('Часовой пояс всех дат', 'Asia/Almaty (UTC+05:00)'),
    ):
        meta.add(label, value)
    for item in LIMITATIONS:
        meta.add('Ограничение', item)
    for item in facts.unavailable_reasons:
        meta.add('Недоступные данные: причина', item)
    if historical_evidence is not None:
        # Explicit allowlist; never flatten arbitrary dictionaries into the file.
        meta.add('Исторические фото', 'НЕТ ПРОВЕРЕННЫХ ИЗОБРАЖЕНИЙ. Синтетическая история не доказывает живое закрытие.')
        for key, label in (
            ('historical_order_count', 'Исторических нарядов'),
            ('historical_submission_count', 'Исторических попыток'),
            ('historical_after_photo_reference_count', 'Исторических ссылок на фото'),
            ('missing_after_photo_row_count', 'Отсутствующих записей фото'),
        ):
            value = historical_evidence.get(key)
            if type(value) is not int or value < 0:
                raise _unavailable()
            meta.add(label, value)

    if selected is None:
        meta.add('Итоги доступны', facts.totals_available)
        if not facts.totals_available:
            meta.add('Итоги смены', 'НЕДОСТАТОЧНО ДАННЫХ: нет полного согласованного снимка и истории. Итоги не вычислены.')
            return title, tables
        metrics = table('Показатели', 'Область', 'Код показателя', 'Показатель', 'Состояние', 'Значение',
                        'Значение точно', 'Числитель', 'Числитель точно', 'Знаменатель',
                        'Подходящих записей', 'Без данных', 'Исключено', 'Малая выборка', 'Таблица источника')
        sources = table('Источники показателей', 'Область', 'Код показателя', 'Когорта', 'ID источника')

        def metric(scope, m):
            metrics.add(scope, m.name, METRICS.get(m.name, m.name), METRIC_STATUS.get(m.status, m.status),
                        m.value, None if m.value is None else str(m.value), m.numerator,
                        None if m.numerator is None else str(m.numerator), m.denominator,
                        m.eligible, m.missing, m.excluded, m.small_sample, m.source_table)
            for label, ids in (('Все подходящие', m.source_ids), ('Без данных', m.missing_source_ids),
                               ('Исключённые', m.excluded_source_ids)):
                for identifier in ids:
                    sources.add(scope, m.name, label, identifier)

        for m in facts.metrics:
            metric('Смена', m)
        ratings = table('Исполнители', 'ID исполнителя', 'Сводный рейтинг', 'Состояние рейтинга')
        for r in facts.ratings:
            ratings.add(r.executor_id, None, 'Недостаточно поддерживаемых исходных показателей')
            for m in (r.human_score, r.closed_on_time, r.closed_with_rework):
                metric(r.executor_id, m)
        materials = table('Материалы закрытых', 'ID материала', 'Название', 'Единица', 'Количество', 'Количество точно')
        material_sources = table('Источники материалов', 'ID материала', 'Тип источника', 'ID источника')
        for m in facts.closed_materials:
            materials.add(m.material_id, m.label, m.unit, m.quantity, str(m.quantity))
            for kind, ids in (('Наряд', m.order_ids), ('Попытка', m.submission_ids), ('Решение', m.review_ids)):
                for identifier in ids:
                    material_sources.add(m.material_id, kind, identifier)
        return title, tables

    order = selected.order
    fields = table('Наряд', 'Поле', 'Значение')
    for label, value in (
        ('ID наряда', order.id), ('Номер', str(order.number)), ('Состояние', STATUS.get(order.status, order.status)),
        ('Тип', 'Плановый' if order.type == 'planned' else 'Внеплановый'), ('Задание', order.description),
        ('Комментарий мастера', order.comment), ('Участок', order.section_id), ('Оборудование', order.equipment_id),
        ('Исполнитель', order.assignment.executor_id), ('Бригада', order.assignment.brigade_id), ('Выдал', order.created_by),
        ('Версия', order.version), ('Ревизия назначения', order.assignment_revision), ('Ревизия срока', order.scheduling_revision),
        ('Выдан', order.issued_at), ('Срок', order.due_at), ('Обновлён', order.updated_at),
        ('Просрочен на время снимка', selected.is_overdue), ('Норма, минут', order.norm_minutes),
        ('Приоритет', {'normal': 'Обычный', 'high': 'Высокий', 'emergency': 'Аварийный'}.get(order.priority, order.priority)),
        ('Текущая попытка', order.current_submission_id),
        ('Сравнение фото до и после', 'Не выполнялось' if order.before_photo_ids else 'Не применимо: фото до работ отсутствуют'),
    ):
        fields.add(label, value)
    attempts = table('Попытки', 'ID попытки', 'Номер', 'Ревизия назначения', 'Текущая', 'Отправил', 'Отправлено',
                     'После срока', 'Полнота', 'Недостающие доказательства', 'Выполненные работы', 'Шифр работы', 'Комментарий')
    reviews = table('Решения мастера', 'ID попытки', 'Решение записано', 'ID решения', 'Мастер', 'Решение',
                    'Оценка мастера', 'Оценка указана', 'Причина', 'Время решения')
    assessments = table('Рекомендации ИИ', 'ID попытки', 'ID оценки', 'Ревизия', 'Режим', 'Модель', 'Версия модели',
                        'Рекомендация', 'Балл ИИ', 'Объяснения', 'Причина резервного режима', 'Устаревшая',
                        'Рекомендация текущей попытки', 'Время оценки', 'Относительно решения мастера',
                        'Версия схемы', 'Длительность, мс')
    ai_evidence = table('Источники ИИ', 'ID оценки', 'ID доказательства')
    photos = table('Ссылки на фото', 'Назначение', 'ID попытки', 'ID фото', 'Состояние')
    materials = table('Материалы попыток', 'ID попытки', 'ID материала', 'Количество', 'Количество точно')
    for identifier in order.before_photo_ids:
        photos.add('До работ', None, identifier, 'Только ссылка; изображение не встроено и не проверялось')
    if not order.before_photo_ids:
        photos.add('До работ', None, None, 'Фото до работ не указаны')
    for attempt in selected.attempts:
        s, review = attempt.submission, attempt.review
        current = s.id == order.current_submission_id and s.assignment_revision == order.assignment_revision
        attempts.add(s.id, s.attempt_number, s.assignment_revision, current, s.submitted_by, s.submitted_at,
                     s.done_late, {'complete': 'Полный', 'incomplete': 'Неполный'}.get(s.completeness, s.completeness),
                     '\n'.join(s.missing_evidence), s.payload.work_description, s.payload.work_code_id, s.payload.comment)
        if review is None:
            reviews.add(s.id, False, None, None, None, None, False, None, None)
        else:
            reviews.add(s.id, True, review.id, review.reviewer_id,
                        'Закрыть' if review.decision == 'close' else 'Вернуть на доработку',
                        review.final_score, review.final_score is not None, review.reason, review.created_at)
        for a in attempt.assessments:
            chronology = ('Нет решения мастера для сравнения' if review is None else
                          'После решения мастера' if a.created_at > review.created_at else
                          'Не позже решения мастера; просмотр не подтверждён')
            assessments.add(s.id, a.id, a.assignment_revision, MODE.get(a.mode, a.mode), a.model, a.model_version,
                            RECOMMENDATION.get(a.recommendation, a.recommendation), a.score, '\n'.join(a.reasons),
                            a.fallback_reason, a.stale, current and not a.stale, a.created_at, chronology,
                            a.schema_version, a.duration_ms)
            for identifier in a.evidence_ids:
                ai_evidence.add(a.id, identifier)
        if not attempt.assessments:
            assessments.add(s.id, None, s.assignment_revision, 'Оценок в источнике нет; состояние задания ИИ неизвестно',
                            *([None] * 12))
        for identifier in s.payload.after_photo_ids:
            photos.add('После работ', s.id, identifier, 'Только ссылка; изображение не встроено и не проверялось')
        if not s.payload.after_photo_ids:
            photos.add('После работ', s.id, None, 'Фото после работ не указаны')
        for m in s.payload.materials:
            materials.add(s.id, m.material_id, m.quantity, str(m.quantity))
    return title, tables


def _shown(value):
    if value is None:
        return 'нет данных'
    if type(value) is bool:
        return 'да' if value else 'нет'
    if isinstance(value, datetime):
        return value.astimezone(ZONE).strftime('%d.%m.%Y %H:%M:%S %z')
    return str(value)


def _xlsx(title, tables, budget):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    book.remove(book.active)
    book.properties.title = title
    book.properties.creator = 'DalaAI'
    book.properties.description = 'Авторизованный снимок. Фото не встроены. Оценки мастера и ИИ разделены.'
    for table in tables:
        budget.check()
        sheet = book.create_sheet(table.title)
        for index, row in enumerate((table.headers, *table.rows), start=1):
            budget.check()
            for col, value in enumerate(row, start=1):
                cell = sheet.cell(index, col)
                if isinstance(value, str):
                    # Explicit string type disables Excel formula/DDE/error inference,
                    # including =, +, -, @, whitespace prefixes and #N/A literals.
                    cell.value = value
                    cell.data_type = 's'
                    cell.quotePrefix = True
                    cell.number_format = '@'
                elif isinstance(value, datetime):
                    cell.value = value.astimezone(ZONE).replace(tzinfo=None)
                    cell.number_format = 'dd.mm.yyyy hh:mm:ss'
                else:
                    cell.value = value
                    if isinstance(value, Decimal):
                        cell.number_format = '0.###############'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                cell.font = Font(name='Calibri', size=11, bold=index == 1,
                                 color='FFFFFF' if index == 1 else '172B3A')
                if index == 1:
                    cell.fill = PatternFill('solid', fgColor='174A61')
            # Long source text remains complete in the cell; desktop row-height
            # limits may require opening the cell. Do not silently truncate it.
            lines = max((sum(max(1, (len(part) + 29) // 30) for part in _shown(v).split('\n')) for v in row), default=1)
            sheet.row_dimensions[index].height = 40 if index == 1 else min(409, max(30, lines * 15))
        for col in range(1, len(table.headers) + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 30 if len(table.headers) > 2 else (42 if col == 1 else 100)
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        sheet.print_title_rows = '1:1'
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
    output = _BoundedBytes()
    book.save(output)
    budget.check()
    return output.getvalue()


def _pdf(title, tables, budget):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen.canvas import Canvas

    for name, path in (('C5Sans', FONT_REGULAR), ('C5SansBold', FONT_BOLD)):
        if not Path(path).is_file():
            raise _unavailable()
        pdfmetrics.registerFont(TTFont(name, path))
    output = _BoundedBytes()
    canvas = Canvas(output, pagesize=A4, pageCompression=1, invariant=1)
    canvas.setTitle(title)
    canvas.setAuthor('DalaAI')
    width, height = A4
    margin, bottom, y, pages = 42, 44, 0, 0
    watermark = str(tables[0].rows[1][1])
    char_widths = {}

    def new_page():
        nonlocal y, pages
        budget.check()
        if pages:
            canvas.showPage()
        pages += 1
        if pages > MAX_PDF_PAGES:
            raise _limit()
        canvas.setFillColorRGB(.09, .29, .38)
        canvas.setFont('C5SansBold', 9)
        canvas.drawString(margin, height - 28, 'DalaAI / НарядAI')
        canvas.setFont('C5Sans', 8)
        canvas.drawRightString(width - margin, height - 28, watermark)
        canvas.setStrokeColorRGB(.7, .77, .8)
        canvas.line(margin, height - 36, width - margin, height - 36)
        canvas.setFont('C5Sans', 8)
        canvas.drawString(margin, 24, 'Asia/Almaty (UTC+05:00) · Фото не встроены')
        canvas.drawRightString(width - margin, 24, 'Страница ' + str(pages))
        y = height - 57

    def text(value, *, bold=False, size=10, gap=3):
        nonlocal y
        budget.check()
        font = 'C5SansBold' if bold else 'C5Sans'
        # Literal canvas text, never parsed as ReportLab XML/HTML or a resource.
        # Character-wise widths yield linear bounded wrapping, even without spaces.
        line, used = '', 0
        lines = []
        for c in _shown(value).replace('\r\n', '\n').replace('\r', '\n').expandtabs(4):
            key = font, size, c
            cw = char_widths.get(key)
            if cw is None:
                cw = pdfmetrics.stringWidth(c, font, size)
                char_widths[key] = cw
            if c == '\n':
                lines.append(line)
                line, used = '', 0
                continue
            if used + cw > width - margin * 2:
                split_at = line.rfind(' ')
                if split_at > 0:
                    lines.append(line[:split_at])
                    line = line[split_at + 1:]
                    used = pdfmetrics.stringWidth(line, font, size)
                else:
                    lines.append(line)
                    line, used = '', 0
            line += c
            used += cw
        lines.append(line)
        for line in lines:
            if y - size * 1.4 < bottom:
                new_page()
            canvas.setFont(font, size)
            canvas.setFillColorRGB(.08, .13, .18)
            canvas.drawString(margin, y, line)
            y -= size * 1.4
        y -= gap

    new_page()
    text(title, bold=True, size=18, gap=12)
    for table in tables:
        if y < bottom + 65:
            new_page()
        text(table.title, bold=True, size=13, gap=7)
        if not table.rows:
            text('Записей нет')
        if table.title.startswith('Источники'):
            # Repeated cohort labels stay once per group in PDF; all source IDs
            # remain present. XLSX keeps normalized, sortable one-ID rows.
            groups = {}
            for row in table.rows:
                groups.setdefault(row[:-1], []).append(_shown(row[-1]))
            for labels, identifiers in groups.items():
                text(' / '.join(_shown(v) for v in labels), bold=True, size=8)
                text(', '.join(identifiers), size=8)
            y -= 8
            continue
        for index, row in enumerate(table.rows, start=1):
            if len(table.headers) == 2:
                text(str(row[0]) + ': ' + _shown(row[1]))
            elif table.title == 'Ссылки на фото':
                text(' / '.join(_shown(v) for v in row), size=8)
            else:
                text('Запись ' + str(index), bold=True)
                for label, value in zip(table.headers, row):
                    text(label + ': ' + _shown(value))
            y -= 4
        y -= 8
    canvas.save()
    budget.check()
    return output.getvalue()


def render_export(facts, output, *, order_id=None, historical_evidence=None):
    """Pure trusted-input renderer; HTTP must use render_export_bounded below."""
    if output not in {'pdf', 'xlsx'}:
        raise ValueError('Unsupported export format')
    budget = _Budget()
    title, tables = _tables(facts, order_id, historical_evidence, budget)
    body = (_pdf if output == 'pdf' else _xlsx)(title, tables, budget)
    if not body or len(body) > MAX_BYTES:
        raise _limit()
    return body


def _render_child():
    """Private subprocess protocol. Stdin is authored only by this module."""
    try:
        import resource
        resource.setrlimit(resource.RLIMIT_CPU, (6, 6))
        resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024, 512 * 1024 * 1024))
        # Never expose this trusted-pickle seam to HTTP or accept artifact files.
        payload = sys.stdin.buffer.read(16 * 1024 * 1024 + 1)
        if len(payload) > 16 * 1024 * 1024:
            raise _limit()
        facts, output, order_id, historical_evidence = pickle.loads(payload)
        body = render_export(facts, output, order_id=order_id, historical_evidence=historical_evidence)
        sys.stdout.buffer.write(b'OK' + body)
    except DomainError as error:
        sys.stdout.buffer.write(b'LIMIT' if error.code == 'REPORT_LIMIT_EXCEEDED' else b'ERROR')
    except Exception:
        sys.stdout.buffer.write(b'ERROR')


def render_export_bounded(facts, output, *, order_id=None, historical_evidence=None):
    """Killable rendering with no inherited session, DB, or credential env."""
    deadline = time.monotonic() + RENDER_SECONDS
    payload = pickle.dumps((facts, output, order_id, historical_evidence), protocol=5)
    if len(payload) > 16 * 1024 * 1024:
        raise _limit()
    # A fresh interpreter receives only import paths and harmless locale settings.
    # Do not inherit DATABASE_URL, provider credentials, cookies or open DB FDs.
    env = {'PATH': os.defpath, 'LANG': 'C.UTF-8', 'PYTHONDONTWRITEBYTECODE': '1',
           'PYTHONPATH': os.pathsep.join(str(Path(p or os.getcwd()).resolve()) for p in sys.path)}
    process = subprocess.Popen([sys.executable, '-m', 'app.reports.c5_exports', '--worker'],
        stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
        env=env, close_fds=True)
    try:
        message, _ = process.communicate(payload, timeout=max(0.001, deadline - time.monotonic()))
        if time.monotonic() > deadline or process.returncode != 0:
            raise _unavailable()
        if message == b'LIMIT':
            raise _limit()
        if not message.startswith(b'OK'):
            raise _unavailable()
        body = message[2:]
        if not body or len(body) > MAX_BYTES:
            raise _limit()
        return body
    except subprocess.TimeoutExpired:
        raise _unavailable() from None
    finally:
        if process.poll() is None:
            process.kill()
        process.communicate()


if __name__ == '__main__' and sys.argv[1:] == ['--worker']:
    _render_child()
