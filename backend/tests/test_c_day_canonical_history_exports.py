"""Full, hash-pinned C1 history through the real six-second C5 renderer.

Run from the repository root:
    PYTHONPATH=backend python -m unittest discover -s backend/tests \
        -p 'test_c_day_canonical_history_exports.py' -v

Source assertions are separate from export acceptance. A renderer timeout is a
failure, never a skip or an expected successful export. This is isolated,
synthetic, offline evidence, not DB/auth, live photo or deployment acceptance.
Only repository code and already-locked renderer dependencies are required.
"""
from collections import Counter, defaultdict
from dataclasses import replace
from datetime import datetime
from decimal import Context, Decimal, localcontext
from functools import lru_cache
import hashlib
import importlib.util
from io import BytesIO
from pathlib import Path
import unittest
from zipfile import ZipFile

from openpyxl import load_workbook

from app.analytics.c3_facts import build_facts
from app.analytics.c3_types import MaterialReference, Period, Provenance, TrustedRows
from app.orders.models import (
    Assignment, Completeness, Decision, DomainError, MaterialUse, MissingEvidence,
    Order, OrderType, Priority, Review, Status, Submission, SubmitPayload,
)
from app.reports import c5_exports as exports


HISTORY_SHA256 = '7d888cdd5bb6a9c01ca7c543fae9e393335d07210dab811aa754f12331d1d2e1'
ROOT = Path(__file__).resolve().parents[2]
SCOPE = ('Isolated complete canonical source scope; Canonical 540 orders,568 '
         'attempts,444 missing photo references; no image bytes')
UNAVAILABLE = 'historical_evidence:unavailable:missing_photo_rows=444'


def instant(value):
    return datetime.fromisoformat(value.replace('Z', '+00:00'))


@lru_cache(maxsize=1)
def canonical_fixture():
    # Load the real standalone generator without changing sys.path or importing
    # the database loader, auth/service modules, or another test's mock fixtures.
    spec = importlib.util.spec_from_file_location(
        'c_day_canonical_generator', ROOT / 'scripts/synthetic/generate.py')
    generator = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(generator)
    history, _ = generator.build_export(seed=20261008, count=540)
    digest = hashlib.sha256(generator.canonical_bytes(history)).hexdigest()
    if digest != HISTORY_SHA256:
        raise AssertionError(f'Canonical history hash changed: {digest}')
    orders = tuple(Order(**dict(
        o, assignment=Assignment(**o['assignment']), status=Status(o['status']),
        type=OrderType(o['type']), priority=Priority(o['priority']),
        issued_at=instant(o['issued_at']), due_at=instant(o['due_at']),
        updated_at=instant(o['updated_at']), before_photo_ids=tuple(o['before_photo_ids']),
    )) for o in history['orders'])
    submissions = tuple(Submission(**dict(
        s, submitted_at=instant(s['submitted_at']), completeness=Completeness(s['completeness']),
        missing_evidence=tuple(MissingEvidence(v) for v in s['missing_evidence']),
        payload=SubmitPayload(**dict(
            s['payload'], after_photo_ids=tuple(s['payload']['after_photo_ids']),
            materials=tuple(MaterialUse(m['material_id'], Decimal(str(m['quantity'])))
                            for m in s['payload']['materials']),
        )),
    )) for s in history['submissions'])
    reviews = tuple(Review(**dict(
        r, decision=Decision(r['decision']), created_at=instant(r['created_at']),
    )) for r in history['reviews'])
    references = tuple(MaterialReference(m['id'], m['label'], m['unit'])
                       for m in history['materials'])
    period = Period(instant('2026-06-30T19:00:00Z'), instant('2026-09-30T19:00:00Z'))
    provenance = Provenance(
        True, 'canonical:' + HISTORY_SHA256, SCOPE, period.end,
        instant('2026-10-08T08:00:00Z'), 'frozen_complete_export', True,
    )
    rows = TrustedRows(provenance, orders, submissions, reviews, (), references)
    facts = build_facts(rows, period)
    facts = replace(facts, unavailable_reasons=facts.unavailable_reasons + (UNAVAILABLE,))
    evidence = {
        'historical_order_count': len(orders),
        'historical_submission_count': len(submissions),
        'historical_after_photo_reference_count': sum(len(s.payload.after_photo_ids) for s in submissions),
        'missing_after_photo_row_count': sum(len(s.payload.after_photo_ids) for s in submissions),
        'physical_evidence_verified': False,
        'historical_completeness_is_verified_evidence': False,
    }
    return history, rows, facts, evidence


def source_tables(facts, evidence):
    budget = exports._Budget()
    _, tables = exports._tables(facts, None, evidence, budget)
    return {table.title: table for table in tables}, budget


def metric_rows(facts):
    yield from (('Смена', metric) for metric in facts.metrics)
    for rating in facts.ratings:
        for metric in (rating.human_score, rating.closed_on_time, rating.closed_with_rework):
            yield rating.executor_id, metric


class CanonicalSourceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.history, cls.rows, cls.facts, cls.evidence = canonical_fixture()

    def test_pinned_full_history_and_unavailable_photo_references(self):
        h = self.history
        self.assertEqual((len(h['orders']), len(h['submissions']), len(h['reviews'])), (540, 568, 568))
        self.assertEqual(Counter(r['decision'] for r in h['reviews']), {'close': 540, 'rework': 28})
        self.assertTrue(h['metadata']['synthetic'])
        self.assertEqual(h['metadata']['photo_evidence'], 'metadata_placeholders_only_no_image_bytes')
        self.assertEqual(h['ai_assessments'], [])
        self.assertEqual(self.rows.assessments, ())
        refs = {identifier: s for s in h['submissions'] for identifier in s['payload']['after_photo_ids']}
        self.assertEqual(len(refs), 444)
        self.assertEqual({p['id'] for p in h['photos']}, set(refs))
        self.assertEqual(len(h['photos']), 444)
        for photo in h['photos']:
            sub = refs[photo['id']]
            self.assertEqual((photo['order_id'], photo['submission_id']), (sub['order_id'], sub['id']))
            self.assertFalse(photo['artifact_available'])
            self.assertEqual(photo['evidence_kind'], 'synthetic_metadata_placeholder')
        self.assertEqual(sum(len(o.attempts) for o in self.facts.orders), 568)
        self.assertTrue(all(a.assessment_status == 'absent' and not a.assessments
                            for o in self.facts.orders for a in o.attempts))

    def assert_mean(self, metric, observations):
        known = [value for _, value in observations if value is not None]
        missing = {identifier for identifier, value in observations if value is None}
        with localcontext(Context(prec=40)):
            numerator = Decimal(sum(known)) if known else None
            expected = (numerator / len(known)).quantize(Decimal('.000000000001')) if known else None
        self.assertEqual((metric.numerator, metric.value, metric.denominator, metric.eligible, metric.missing),
                         (numerator, expected, len(known), len(observations), len(missing)))
        self.assertEqual(set(metric.source_ids), {identifier for identifier, _ in observations})
        self.assertEqual(set(metric.missing_source_ids), missing)
        self.assertEqual(metric.status, 'no_cohort' if not observations else 'missing' if not known
                         else 'partial' if missing else 'ok')

    def test_totals_denominators_nulls_and_sources_from_canonical_rows(self):
        h = self.history
        subs = {s['id']: s for s in h['submissions']}
        closes = [r for r in h['reviews'] if r['decision'] == 'close']
        reworked = {subs[r['submission_id']]['order_id'] for r in h['reviews'] if r['decision'] == 'rework'}
        counts = {
            'issued_orders': {o['id'] for o in h['orders']},
            'submitted_orders': {s['order_id'] for s in h['submissions']},
            'submission_attempts': set(subs),
            'closed_orders': {subs[r['submission_id']]['order_id'] for r in closes},
            'rework_decisions': {r['id'] for r in h['reviews'] if r['decision'] == 'rework'},
            'awaiting_review': set(), 'overdue_active': set(),
        }
        metrics = {m.name: m for m in self.facts.metrics}
        self.assertEqual(set(metrics), set(counts) | {'human_score', 'closed_on_time', 'closed_with_rework', 'attempt_on_time'})
        for name, identifiers in counts.items():
            metric = metrics[name]
            self.assertEqual((metric.value, metric.numerator, metric.denominator),
                             (Decimal(len(identifiers)), Decimal(len(identifiers)), None))
            self.assertEqual((metric.eligible, metric.missing, metric.status), (len(identifiers), 0, 'ok'))
            self.assertEqual(set(metric.source_ids), identifiers)
        self.assert_mean(metrics['attempt_on_time'], [(s['id'], int(not s['done_late'])) for s in subs.values()])
        by_executor = defaultdict(list)
        for review in closes:
            by_executor[subs[review['submission_id']]['submitted_by']].append(review)
        self.assertEqual({r.executor_id for r in self.facts.ratings}, set(by_executor))
        scopes = [(metrics, closes)] + [(
            {m.name: m for m in (r.human_score, r.closed_on_time, r.closed_with_rework)}, by_executor[r.executor_id],
        ) for r in self.facts.ratings]
        for components, cohort in scopes:
            self.assert_mean(components['human_score'], [(r['id'], r['final_score']) for r in cohort])
            self.assert_mean(components['closed_on_time'],
                             [(r['submission_id'], int(not subs[r['submission_id']]['done_late'])) for r in cohort])
            self.assert_mean(components['closed_with_rework'],
                             [(subs[r['submission_id']]['order_id'], int(subs[r['submission_id']]['order_id'] in reworked))
                              for r in cohort])
        self.assertEqual((metrics['human_score'].denominator, metrics['human_score'].missing), (486, 54))
        self.assertEqual(metrics['human_score'].value, Decimal('86.129629629630'))
        self.assertTrue(all(r.composite_score is None and r.composite_status == 'unsupported_inputs'
                            for r in self.facts.ratings))
        quantities, material_sources = defaultdict(Decimal), defaultdict(set)
        for review in closes:
            sub = subs[review['submission_id']]
            for use in sub['payload']['materials']:
                quantities[use['material_id']] += Decimal(str(use['quantity']))
                material_sources[use['material_id']].add((sub['order_id'], sub['id'], review['id']))
        self.assertEqual({m.material_id: m.quantity for m in self.facts.closed_materials}, dict(quantities))
        for material in self.facts.closed_materials:
            source = material_sources[material.material_id]
            self.assertEqual((set(material.order_ids), set(material.submission_ids), set(material.review_ids)),
                             tuple({row[i] for row in source} for i in range(3)))

    def test_full_source_tables_provenance_and_unchanged_caps(self):
        self.assertEqual((exports.RENDER_SECONDS, exports.MAX_BYTES, exports.MAX_CELLS,
                          exports.MAX_TEXT_CHARS, exports.MAX_CELL_CHARS, exports.MAX_PDF_PAGES),
                         (6.0, 8 * 1024 * 1024, 120_000, 1_000_000, 30_000, 200))
        tables, budget = source_tables(self.facts, self.evidence)
        self.assertEqual((budget.cells, budget.chars), (30_516, 617_427))
        meta = dict(tables['Об отчёте'].rows)
        self.assertEqual(meta['Источник'], 'canonical:' + HISTORY_SHA256)
        self.assertEqual(meta['Разрешённая область'], SCOPE)
        self.assertEqual(meta['Режим'], 'СИНТЕТИЧЕСКИЕ ДАННЫЕ')
        self.assertEqual(meta['Полнота выборки'], 'frozen_complete_export')
        self.assertTrue(meta['История полная по заявлению источника'])
        self.assertTrue(meta['Итоги доступны'])
        self.assertIn('НЕТ ПРОВЕРЕННЫХ ИЗОБРАЖЕНИЙ', meta['Исторические фото'])
        self.assertEqual([meta[key] for key in ('Исторических нарядов', 'Исторических попыток',
                                              'Исторических ссылок на фото', 'Отсутствующих записей фото')],
                         [540, 568, 444, 444])
        self.assertIn(('Недоступные данные: причина', UNAVAILABLE), tables['Об отчёте'].rows)
        actual = {(row[0], row[1]): row for row in tables['Показатели'].rows}
        expected_sources = []
        for scope, metric in metric_rows(self.facts):
            row = actual[scope, metric.name]
            self.assertEqual(row[4:12], (metric.value, str(metric.value), metric.numerator,
                                        str(metric.numerator), metric.denominator, metric.eligible,
                                        metric.missing, metric.excluded))
            for label, ids in (('Все подходящие', metric.source_ids), ('Без данных', metric.missing_source_ids),
                               ('Исключённые', metric.excluded_source_ids)):
                expected_sources.extend((scope, metric.name, label, identifier) for identifier in ids)
        self.assertEqual(Counter(tables['Источники показателей'].rows), Counter(expected_sources))
        self.assertEqual(len(tables['Исполнители'].rows), 15)
        self.assertTrue(all(row[1] is None for row in tables['Исполнители'].rows))


class CanonicalExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        _, _, cls.facts, cls.evidence = canonical_fixture()

    def render(self, output, **kwargs):
        try:
            return exports.render_export_bounded(self.facts, output, historical_evidence=self.evidence, **kwargs)
        except DomainError as exc:
            self.fail(f'Canonical {output.upper()} export failed: {exc.code}; '
                      'the real six-second renderer and output caps remain unchanged')

    def assert_safe_workbook(self, body):
        self.assertLessEqual(len(body), exports.MAX_BYTES)
        book = load_workbook(BytesIO(body), data_only=False)
        for sheet in book:
            for row in sheet:
                for cell in row:
                    self.assertNotIn(cell.data_type, {'f', 'e'})
                    self.assertIsNone(cell.hyperlink)
                    if isinstance(cell.value, str):
                        self.assertEqual(cell.data_type, 's')
                        self.assertTrue(cell.quotePrefix)
        with ZipFile(BytesIO(body)) as archive:
            self.assertFalse(any('vbaProject' in name or 'externalLinks' in name or 'media/' in name
                                 for name in archive.namelist()))
            for name in archive.namelist():
                if name.endswith('.rels'):
                    self.assertNotIn(b'TargetMode="External"', archive.read(name))
        return book

    def test_full_canonical_xlsx_with_real_bounded_renderer(self):
        body = self.render('xlsx')
        book = self.assert_safe_workbook(body)
        tables, _ = source_tables(self.facts, self.evidence)
        self.assertEqual(book.sheetnames, list(tables))
        self.assertEqual(sum(sheet.max_row * sheet.max_column for sheet in book), 30_516)
        for name, table in tables.items():
            actual = list(book[name].values)
            self.assertEqual(actual[0], table.headers)
            self.assertEqual(len(actual), len(table.rows) + 1)
            for actual_row, expected_row in zip(actual[1:], table.rows):
                for value, expected in zip(actual_row, expected_row):
                    if isinstance(expected, Decimal):
                        self.assertIsInstance(value, (int, float))
                        self.assertAlmostEqual(value, float(expected), places=12)
                    elif isinstance(expected, datetime):
                        self.assertEqual(value, expected.astimezone(exports.ZONE).replace(tzinfo=None))
                    else:
                        self.assertEqual(value, expected)

    def test_full_canonical_pdf_with_real_bounded_renderer(self):
        body = self.render('pdf')
        self.assertLessEqual(len(body), exports.MAX_BYTES)
        self.assertTrue(body.startswith(b'%PDF-1.'))
        self.assertTrue(body.rstrip().endswith(b'%%EOF'))
        self.assertIn(b'/FontFile2', body)
        self.assertIn(b'/ToUnicode', body)
        for forbidden in (b'/JavaScript', b'/OpenAction', b'/URI', b'/EmbeddedFile', b'/Subtype /Image'):
            self.assertNotIn(forbidden, body)
        # Content/provenance is checked independently above. Text extraction and
        # visual PDF acceptance require separate evidence, not a new dependency.

    def test_bounded_literal_canary_and_null_score_derived_from_history(self):
        original = next(o for o in self.facts.orders if o.attempts[-1].review.final_score is None)
        attacks = ('=HYPERLINK("https://example.invalid","x")', '\r\t@SUM(A1:A2)',
                   '#N/A', '<img src="https://example.invalid/a"> & _x000D_\rtext')
        attempt = original.attempts[-1]
        changed = replace(original, order=replace(original.order, description=attacks[0], comment=attacks[1]),
                          attempts=(*original.attempts[:-1], replace(attempt, submission=replace(
                              attempt.submission, payload=replace(attempt.submission.payload,
                                                                  work_description=attacks[2], comment=attacks[3])))))
        # Explicitly label this small mutation: it is not the pinned canonical export.
        facts = replace(self.facts, orders=(changed,), provenance=replace(
            self.facts.provenance, source_ref='synthetic-literal-canary:derived-from:' + HISTORY_SHA256,
            scope_description='Single-order literal-text canary derived from synthetic history'))
        body = exports.render_export_bounded(facts, 'xlsx', order_id=changed.order.id)
        book = self.assert_safe_workbook(body)
        values = [cell.value for sheet in book for row in sheet for cell in row]
        for attack in attacks:
            self.assertIn(attack, values)
        review_row = next(row for row in book['Решения мастера'].values if row[0] == attempt.submission.id)
        self.assertIsNone(review_row[5])
        self.assertFalse(review_row[6])
        self.assertNotIn('Показатели', book.sheetnames)


if __name__ == '__main__':
    unittest.main()
