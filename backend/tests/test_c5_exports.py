"""C5 file round-trips + HTTP tests using C111's real auth and mock DB transport.

Fixtures are synthetic. These tests are not PostgreSQL, production load or
native Microsoft Excel acceptance. No auth shortcuts are implemented here.
"""
from dataclasses import replace
from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO
import json
import os
from pathlib import Path
import subprocess
import sys
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.analytics.c3_repository import RuntimeReportService
from app.core.auth_boundary import SESSION_COOKIE_NAME
from app.orders.models import DomainError
from app.reports import c5_exports as exports
from app.reports.c4_routes import create_c_runtime_router
from app.reports.c5_export_routes import create_c_export_router
from app.sessions.service import SessionService
from test_c4_render import fixture, uid as fact_uid, ATTACK
from test_c_runtime_routes import MockDatabase, Clock, NOW, uid, canonical_history_fixture


def workbook(body):
    return load_workbook(BytesIO(body), data_only=False)


def cells(book):
    return [c for sheet in book for row in sheet for c in row]


class ExportFileTests(unittest.TestCase):
    def setUp(self):
        self.facts = fixture()
        self.order_id = fact_uid(1)

    def render(self, output='xlsx', **kwargs):
        return exports.render_export(self.facts, output, order_id=self.order_id, **kwargs)

    def test_xlsx_numbers_dates_null_zero_and_human_ai_are_distinct(self):
        book = workbook(self.render())
        reviews, ai, materials = book['Решения мастера'], book['Рекомендации ИИ'], book['Материалы попыток']
        self.assertIsNone(reviews['F3'].value)  # Human null never uses AI 91.
        self.assertEqual(reviews['G3'].value, False)
        self.assertEqual(ai['H2'].value, 91)
        self.assertEqual(ai['H2'].data_type, 'n')
        self.assertEqual(ai['L2'].value, False)  # Old/stale recommendation.
        self.assertIsNone(ai['H3'].value)
        self.assertEqual(materials['C2'].value, 1.125)
        self.assertEqual(materials['C2'].data_type, 'n')
        self.assertEqual(materials['D2'].value, '1.125')
        self.assertEqual(materials['D2'].data_type, 's')
        self.assertIsInstance(reviews['I3'].value, datetime)
        self.assertIsNone(reviews['I3'].value.tzinfo)
        self.assertEqual(reviews['I3'].value.hour, 11)  # 06:55 UTC -> 11:55 Almaty.
        current = self.facts.orders[0].attempts[1]
        changed = replace(current, review=replace(current.review, final_score=0))
        self.facts = replace(self.facts, orders=(replace(self.facts.orders[0], attempts=(changed,)),))
        book = workbook(self.render())
        self.assertEqual(book['Решения мастера']['F2'].value, 0)
        self.assertEqual(book['Решения мастера']['F2'].data_type, 'n')
        self.assertTrue(book['Решения мастера']['G2'].value)

    def test_formula_dde_hyperlink_error_and_xml_text_are_literal(self):
        attacks = ['=HYPERLINK("https://evil.invalid","x")', '+cmd|\'/C calc\'!A0',
                   '-1+2', '@SUM(A1:A2)', '\t=1+1', '\r=1+1', '\r\n=1+1', 'before\rafter\nend',
                   '\r\t@SUM(A1:A2)', 'literal &#13; and _x000D_', '_x005F_', '_x005F_x000D_',
                   '_x000D_\r_x005F_', '#N/A', ATTACK]
        for attack in attacks:
            with self.subTest(attack=attack):
                order = replace(self.facts.orders[0].order, description=attack)
                facts = replace(self.facts, orders=(replace(self.facts.orders[0], order=order),))
                body = exports.render_export(facts, 'xlsx', order_id=self.order_id)
                book = workbook(body)
                cell = next(c for c in cells(book) if c.value == attack)
                self.assertEqual(cell.data_type, 's')
                self.assertTrue(cell.quotePrefix)
                self.assertIsNone(cell.hyperlink)
                self.assertFalse(any(c.data_type in {'f', 'e'} for c in cells(book)))
                self.assertEqual(book['Материалы попыток']['C2'].value, 1.125)
                self.assertEqual(book['Материалы попыток']['C2'].data_type, 'n')
                self.assertIsInstance(book['Решения мастера']['I3'].value, datetime)
                self.assertIsNone(book['Решения мастера']['F3'].value)
                with ZipFile(BytesIO(body)) as archive:
                    for name in archive.namelist():
                        if name.endswith('.rels'):
                            self.assertNotIn(b'TargetMode="External"', archive.read(name))
                    self.assertFalse(any('vbaProject' in name or 'externalLinks' in name or 'media/' in name
                                         for name in archive.namelist()))

    def test_literal_roundtrip_in_fresh_process_with_each_xml_backend(self):
        # Explicitly disable lxml in a fresh interpreter. Patching openpyxl.LXML
        # after import does not switch its already-bound XML writer functions.
        script = r"""
import importlib.util
import json
import os
import unittest
import openpyxl
from test_c5_exports import ExportFileTests
expected = os.environ['OPENPYXL_LXML'] == 'True' and importlib.util.find_spec('lxml') is not None
assert openpyxl.LXML is expected, (openpyxl.LXML, expected)
suite = unittest.TestSuite(ExportFileTests(name) for name in (
    'test_formula_dde_hyperlink_error_and_xml_text_are_literal',
    'test_xlsx_numbers_dates_null_zero_and_human_ai_are_distinct',
))
result = unittest.TextTestRunner().run(suite)
assert result.wasSuccessful()
print(json.dumps({'lxml': openpyxl.LXML, 'tests': result.testsRun}))
"""
        for mode in ('False', 'True'):
            with self.subTest(mode=mode):
                env = {**os.environ, 'OPENPYXL_LXML': mode,
                       'PYTHONPATH': os.pathsep.join(str(Path(p or os.getcwd()).resolve()) for p in sys.path)}
                result = subprocess.run([sys.executable, '-c', script], env=env,
                                        capture_output=True, text=True, timeout=15, check=False)
                self.assertEqual(result.returncode, 0, result.stderr)
                proof = json.loads(result.stdout)
                self.assertEqual(proof['tests'], 2)
                if mode == 'False':
                    self.assertFalse(proof['lxml'])

    def test_long_text_is_not_silently_truncated_and_order_is_scoped(self):
        other = replace(self.facts.orders[0], order=replace(self.facts.orders[0].order,
                         id=fact_uid(999), description='FOREIGN-ORDER-CANARY'))
        self.facts = replace(self.facts, orders=self.facts.orders + (other,))
        body = self.render()
        book = workbook(body)
        self.assertEqual(book['Попытки']['J2'].value,
                         self.facts.orders[0].attempts[0].submission.payload.work_description)
        self.assertEqual(len(book['Попытки']['J2'].value), 6000)
        self.assertNotIn('FOREIGN-ORDER-CANARY', str([c.value for c in cells(book)]))
        self.assertNotIn('Показатели', book.sheetnames)
        for unknown in (fact_uid(888), ''):
            with self.assertRaises(ValueError):
                exports.render_export(self.facts, 'xlsx', order_id=unknown)
        duplicate = replace(self.facts, orders=(self.facts.orders[0],)*2)
        with self.assertRaises(ValueError):
            exports.render_export(duplicate, 'pdf', order_id=self.order_id)

    def test_shift_incomplete_totals_are_unavailable_never_zero(self):
        facts = replace(self.facts, provenance=replace(self.facts.provenance, coverage='partial_keyset'))
        body = exports.render_export(facts, 'xlsx')
        book = workbook(body)
        self.assertEqual(book.sheetnames, ['Об отчёте'])
        values = dict(book['Об отчёте'].values)
        self.assertFalse(values['Итоги доступны'])
        self.assertIn('НЕДОСТАТОЧНО ДАННЫХ', values['Итоги смены'])

    def test_shift_uses_typed_fact_values_exact_decimal_and_sources(self):
        facts = replace(self.facts, metrics=(replace(self.facts.metrics[0],
                        value=Decimal('0.123456789123456789')),))
        body = exports.render_export(facts, 'xlsx')
        book = workbook(body)
        self.assertEqual(book['Показатели']['F2'].value, '0.123456789123456789')
        self.assertEqual(book['Показатели']['E2'].data_type, 'n')
        self.assertIn(fact_uid(31), [row[3] for row in book['Источники показателей'].values])
        self.assertIsNone(book['Исполнители']['B2'].value)
        self.assertEqual(book['Материалы закрытых']['D2'].value, 1.125)

    def test_historical_missing_photo_disclosure_is_visible_and_allowlisted(self):
        history = dict(historical_order_count=1, historical_submission_count=2,
                       historical_after_photo_reference_count=2, missing_after_photo_row_count=2,
                       session_token='SECRET-CANARY', arbitrary_url='https://secret.invalid')
        book = workbook(self.render(historical_evidence=history))
        values = dict(book['Об отчёте'].values)
        self.assertIn('НЕТ ПРОВЕРЕННЫХ ИЗОБРАЖЕНИЙ', values['Исторические фото'])
        self.assertEqual(values['Отсутствующих записей фото'], 2)
        self.assertNotIn('SECRET-CANARY', str([c.value for c in cells(book)]))
        self.assertIn('Фото до работ не указаны', [row[3] for row in book['Ссылки на фото'].values])

    def test_pdf_is_real_embedded_cyrillic_and_has_no_active_resources(self):
        body = self.render('pdf')
        self.assertTrue(body.startswith(b'%PDF-1.'))
        self.assertTrue(body.rstrip().endswith(b'%%EOF'))
        self.assertIn(b'/FontFile2', body)
        self.assertIn(b'/ToUnicode', body)
        for forbidden in (b'/JavaScript', b'/OpenAction', b'/URI', b'/EmbeddedFile', b'/Subtype /Image'):
            self.assertNotIn(forbidden, body)
        # pypdf/Poppler extraction + raster inspection are separate evidence;
        # no additional runtime dependency is introduced just for this check.

    def test_limits_font_failure_control_chars_and_no_plain_dict(self):
        for name in ('MAX_BYTES', 'MAX_CELLS', 'MAX_TEXT_CHARS', 'MAX_CELL_CHARS', 'MAX_PDF_PAGES'):
            output = 'pdf' if name == 'MAX_PDF_PAGES' else 'xlsx'
            with self.subTest(name=name), patch.object(exports, name, 1):
                with self.assertRaises(DomainError) as caught:
                    self.render(output)
                self.assertEqual(caught.exception.code, 'REPORT_LIMIT_EXCEEDED')
        with patch.object(exports, 'FONT_REGULAR', '/missing/font.ttf'):
            with self.assertRaises(DomainError):
                self.render('pdf')
        with self.assertRaises(TypeError):
            exports.render_export({}, 'xlsx')
        order = replace(self.facts.orders[0].order, description='bad\x00text')
        facts = replace(self.facts, orders=(replace(self.facts.orders[0], order=order),))
        with self.assertRaises(DomainError):
            exports.render_export(facts, 'xlsx', order_id=self.order_id)

    def test_real_bounded_subprocess_both_formats(self):
        for output in ('pdf', 'xlsx'):
            body = exports.render_export_bounded(self.facts, output, order_id=self.order_id)
            self.assertTrue(body.startswith(b'%PDF' if output == 'pdf' else b'PK'))
            self.assertLess(len(body), exports.MAX_BYTES)

    def test_subprocess_timeout_kills_child_and_environment_has_no_secrets(self):
        from unittest.mock import MagicMock
        child = MagicMock()
        child.communicate.side_effect = [subprocess.TimeoutExpired('renderer', 6), (b'', None)]
        child.poll.return_value = None
        with patch.dict(os.environ, {'DATABASE_URL': 'SECRET', 'AI_API_KEY': 'SECRET'}), \
             patch.object(exports.subprocess, 'Popen', return_value=child) as spawn:
            with self.assertRaises(DomainError) as caught:
                self.render_bounded()
            self.assertEqual(caught.exception.code, 'TEMPORARILY_UNAVAILABLE')
            child.kill.assert_called_once()
            self.assertNotIn('DATABASE_URL', spawn.call_args.kwargs['env'])
            self.assertNotIn('AI_API_KEY', spawn.call_args.kwargs['env'])
            self.assertTrue(spawn.call_args.kwargs['close_fds'])

    def render_bounded(self):
        return exports.render_export_bounded(self.facts, 'xlsx', order_id=self.order_id)


class ExportRouteTests(unittest.TestCase):
    def setUp(self):
        self.db = MockDatabase()
        self.real, self.domain = Clock(), Clock()
        sessions = SessionService(lambda: self.db, allowed_origin='https://reports.test', real_clock=self.real)
        service = RuntimeReportService(sessions, domain_clock=self.domain, synthetic=True)
        app = FastAPI()
        # Integration order is part of the contract: UUID.pdf is not C4's UUID.
        app.include_router(create_c_export_router(service))
        app.include_router(create_c_runtime_router(service))
        self.client = TestClient(app)
        self.headers = {'cookie': SESSION_COOKIE_NAME + '=' + self.db.handle}
        self.query = {'start': (NOW-timedelta(days=1)).isoformat(), 'end': NOW.isoformat()}
        self.patch = patch('app.reports.c5_export_routes.render_export_bounded', side_effect=exports.render_export)
        self.renderer = self.patch.start()
        self.addCleanup(self.patch.stop)

    def get(self, path='/reports/shift.xlsx', query=None, headers=None):
        return self.client.get('/api/v1' + path, params=self.query if query is None else query,
                               headers=self.headers if headers is None else headers)

    def test_four_fixed_routes_and_cache_safe_download_names(self):
        for kind in ('shift', 'orders/' + uid(10)):
            for ext in ('pdf', 'xlsx'):
                with self.subTest(kind=kind, ext=ext):
                    response = self.get('/reports/' + kind + '.' + ext)
                    self.assertEqual(response.status_code, 200, response.text[:100] if ext == 'pdf' else '')
                    self.assertTrue(response.content.startswith(b'%PDF' if ext == 'pdf' else b'PK'))
                    self.assertIn('attachment; filename="naryadai-', response.headers['content-disposition'])
                    self.assertIn('private, no-store', response.headers['cache-control'])
                    self.assertEqual(response.headers['vary'], 'Cookie')
                    self.assertEqual(response.headers['x-content-type-options'], 'nosniff')
                    self.assertEqual(response.headers['referrer-policy'], 'no-referrer')
                    self.assertEqual(int(response.headers['content-length']), len(response.content))
        self.assertEqual(self.get('/reports/orders/' + uid(10)).status_code, 200)

    def test_session_only_missing_duplicate_expired_revoked(self):
        for headers in ({}, {'Authorization': 'Bearer ' + self.db.handle},
                        {'X-Actor-Id': uid(3), 'X-Role': 'master'},
                        {'cookie': self.headers['cookie'] + '; ' + self.headers['cookie']}):
            self.assertEqual(self.get(headers=headers).status_code, 401)
        self.db.expiry = NOW
        self.assertEqual(self.get().status_code, 401)
        self.db.expiry = NOW + timedelta(hours=1)
        self.db.revoked = NOW
        self.assertEqual(self.get().status_code, 401)
        self.renderer.assert_not_called()

    def test_current_role_scope_and_foreign_order_deny_before_render(self):
        for role in ('manager', 'executor', 'admin'):
            self.db.role = role
            self.assertEqual(self.get().status_code, 403)
        self.db.role = 'master'
        self.db.active = False
        self.assertIn(self.get().status_code, (401, 403))
        self.db.active = True
        self.db.sections = []
        self.assertEqual(self.get().status_code, 403)
        self.db.sections = [uid(999)]
        self.assertEqual(self.get('/reports/orders/' + uid(10) + '.xlsx').status_code, 404)
        self.renderer.assert_not_called()

    def test_existing_strict_query_and_uuid_validation(self):
        cases = [({}, 422), ({**self.query, 'format': 'pdf'}, 422),
                 ({**self.query, 'format': 'html'}, 422), ({**self.query, 'token': self.db.handle}, 400),
                 ({**self.query, 'role': 'master'}, 400), ({**self.query, 'section_id': uid(999)}, 400),
                 ([*self.query.items(), ('start', self.query['start'])], 400),
                 ({**self.query, 'end': (NOW+timedelta(seconds=1)).isoformat()}, 422),
                 ({**self.query, 'start': (NOW-timedelta(days=94)).isoformat()}, 422)]
        for query, status in cases:
            with self.subTest(query=query):
                self.assertEqual(self.get(query=query).status_code, status)
        self.assertEqual(self.get('/reports/orders/not-an-id.pdf').status_code, 422)
        self.renderer.assert_not_called()

    def test_expiry_and_role_change_during_serialization_never_release_body(self):
        for mutation, status in ((lambda: setattr(self.real, 'value', self.db.expiry), 401),
                                 (lambda: setattr(self.db, 'role', 'manager'), 403),
                                 (lambda: setattr(self.db, 'sections', [uid(999)]), 403)):
            self.real.value, self.db.role, self.db.sections = NOW, 'master', [uid(1)]
            def render(*args, **kwargs):
                body = exports.render_export(*args, **kwargs)
                mutation()
                return body
            self.renderer.side_effect = render
            response = self.get()
            self.assertEqual(response.status_code, status)
            self.assertNotIn('content-disposition', response.headers)
            self.assertTrue(response.headers['content-type'].startswith('application/json'))

    def test_synthetic_history_missing_photos_are_visible(self):
        identifier, _, _, _ = canonical_history_fixture(self.db)
        response = self.get('/reports/orders/' + identifier + '.xlsx')
        self.assertEqual(response.status_code, 200, response.text[:200] if response.status_code != 200 else '')
        book = workbook(response.content)
        meta = dict(book['Об отчёте'].values)
        self.assertIn('НЕТ ПРОВЕРЕННЫХ ИЗОБРАЖЕНИЙ', meta['Исторические фото'])
        self.assertEqual(meta['Отсутствующих записей фото'], 1)
        self.assertNotIn('Изображение проверено', str(list(book['Ссылки на фото'].values)))

    def test_renderer_errors_size_and_capacity_fail_closed(self):
        for failure, status in ((RuntimeError('SECRET-DB-DETAILS'), 503),
                                (DomainError('REPORT_LIMIT_EXCEEDED', 'Too large'), 422)):
            self.renderer.side_effect = failure
            response = self.get()
            self.assertEqual(response.status_code, status)
            self.assertNotIn('SECRET-DB-DETAILS', response.text)
            self.assertNotIn('content-disposition', response.headers)
        self.renderer.side_effect = None
        self.renderer.return_value = b'x' * (exports.MAX_BYTES+1)
        self.assertEqual(self.get().status_code, 422)
        with patch('app.reports.c5_export_routes._SLOTS') as slots:
            slots.acquire.return_value = False
            response = self.get()
            self.assertEqual(response.status_code, 503)
            self.assertEqual(response.headers['retry-after'], '1')
            slots.release.assert_not_called()


if __name__ == '__main__':
    unittest.main()
